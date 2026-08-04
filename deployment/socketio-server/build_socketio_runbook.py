from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("SukalyanAI_SocketIO_Deployment_Runbook.docx")

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
CYAN = "00A6C8"
INK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_CYAN = "EAF9FC"
WHITE = "FFFFFF"
GREEN = "16794A"
AMBER = "8A5B00"
RED = "9B1C1C"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "left" if side == "start" else "right" if side == "end" else side
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=11, color=INK, bold=False, italic=False, font="Calibri") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, text, end])


def add_text(doc, text: str, *, bold=False, italic=False, color=INK, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    for run in p.runs:
        set_run(run)
    return p


def add_step(doc, title: str, detail: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(title + " "), bold=True, color=DARK_BLUE)
    set_run(p.add_run(detail))
    return p


def add_code(doc, lines: str):
    for line in lines.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.10)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F6F8FA")
        pPr.append(shd)
        set_run(p.add_run(line if line else " "), size=9, color=NAVY, font="Consolas")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_callout(doc, label: str, text: str, fill=PALE_CYAN, accent=CYAN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(label + "  "), size=10, bold=True, color=accent)
    set_run(p.add_run(text), size=10, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(header), size=9.5, bold=True, color=NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            set_run(p.add_run(value), size=9.5, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle.paragraph_format.space_after = Pt(18)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section, *, first=False) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = first


def configure_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("SUKALYANAI  /  SOCKET.IO ALERT SYSTEM"), size=8.5, bold=True, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("Deployment Runbook  |  "), size=8.5, color=MUTED)
    add_field(p, "PAGE")


def add_heading(doc, level: int, text: str):
    return doc.add_heading(text, level=level)


def build() -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_section(section, first=True)
    configure_header_footer(section)

    # Editorial-cover pattern with restrained technical styling.
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    set_run(p.add_run("DEPLOYMENT & OPERATIONS RUNBOOK"), size=10, bold=True, color=CYAN)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("SukalyanAI Socket.IO\nAlert Delivery System")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Contabo relay, VisualAI Edge publisher, and Hostinger frontend")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    set_run(p.add_run("Production reference • Version 1.0 • 29 July 2026"), size=10.5, color=MUTED)
    add_callout(
        doc,
        "PURPOSE",
        "A repeatable installation and support guide for sending VisualAI alerts "
        "from the factory edge computer to the public SukalyanAI website.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    set_run(p.add_run("Archive classification: Internal technical operations"), size=9.5, color=MUTED, italic=True)
    doc.add_page_break()

    add_heading(doc, 1, "1. System at a glance")
    add_callout(
        doc,
        "LIVE FLOW",
        "RTSP camera + VisualAI Edge  ->  HTTPS/WSS  ->  Contabo Socket.IO relay  "
        "->  sukalyanai.com browser alert table",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    add_table(
        doc,
        ["Component", "Location", "Responsibility"],
        [
            ["Camera / AI processing", "Home or factory edge PC", "Reads RTSP, detects people, builds alert JSON."],
            ["Alert publisher", "VisualAI-Edge / Projects / Jutemill", "Emits the alert to Contabo using Socket.IO authentication."],
            ["Public relay", "Contabo VPS", "Accepts authenticated publishers and broadcasts alerts to browsers."],
            ["Reverse proxy + TLS", "Contabo Nginx / Let's Encrypt", "Exposes secure HTTPS/WSS without opening port 3000."],
            ["Alert display", "Hostinger website", "Receives alert_received and inserts newest records first."],
        ],
        [2100, 2500, 4760],
    )
    add_text(
        doc,
        "Important separation: Hostinger serves the web page; Contabo carries the live Socket.IO connection. "
        "The edge PC makes an outbound connection, so no static public IP or inbound port-forwarding is needed for alerts.",
    )

    add_heading(doc, 1, "2. Network and prerequisites")
    add_bullet(doc, "Debian or Ubuntu Contabo VPS with root/sudo access.")
    add_bullet(doc, "DNS A record: socket.sukalyanai.com points to the Contabo public IPv4 address.")
    add_bullet(doc, "TCP 22 (or the configured SSH port), 80, and 443 allowed. Port 3000 remains bound to 127.0.0.1.")
    add_bullet(doc, "A valid email address for Let's Encrypt expiry notices.")
    add_bullet(doc, "Browser origins: https://www.sukalyanai.com and https://sukalyanai.com.")
    add_callout(
        doc,
        "DNS CHECK",
        "Before installation, verify that nslookup socket.sukalyanai.com returns the Contabo IP. "
        "Certificate creation will fail if DNS is not ready.",
        fill="FFF6E5",
        accent=AMBER,
    )

    add_heading(doc, 1, "3. One-command Contabo installation")
    add_step(doc, "Copy the installer.", "From Windows PowerShell, upload it to the VPS.")
    add_code(
        doc,
        r'''scp "D:\SukalyanAI\Industrial-AI\VisualAI-Edge\deployment\socketio-server\install_socketio_contabo.sh" root@YOUR_VPS_IP:/root/''',
    )
    add_step(doc, "Run it on Contabo.", "Replace the certificate email address before executing.")
    add_code(
        doc,
        """chmod +x /root/install_socketio_contabo.sh
sudo /root/install_socketio_contabo.sh \\
  --domain socket.sukalyanai.com \\
  --email YOUR_LETS_ENCRYPT_EMAIL \\
  --origins https://www.sukalyanai.com,https://sukalyanai.com""",
    )
    add_step(doc, "Save the generated token.", "Copy the final SOCKET_AUTH_TOKEN value into a password manager.")
    add_callout(
        doc,
        "SAFE TO RERUN",
        "The installer backs up the current server, environment, systemd, and Nginx files. "
        "It reuses the existing token unless --rotate-token is passed.",
    )
    add_text(doc, "For an HTTP-only staging pass before DNS is ready, add --skip-certbot. Rerun without that option after DNS resolves.")

    add_heading(doc, 1, "4. What the installer creates")
    add_table(
        doc,
        ["Path", "Purpose", "Ownership / mode"],
        [
            ["/opt/socketio-server/server.js", "Socket.IO relay application", "socketio:socketio / 0640"],
            ["/opt/socketio-server/package.json", "Pinned runtime dependency definition", "socketio:socketio / 0640"],
            ["/etc/socketio-server.env", "Host, port, origins, and secret token", "root:socketio / 0640"],
            ["/etc/systemd/system/socketio-server.service", "Hardened auto-restarting service", "root-owned"],
            ["/etc/nginx/sites-available/socketio-server", "TLS reverse proxy and WebSocket upgrade", "root-owned"],
        ],
        [3300, 3860, 2200],
    )
    add_text(
        doc,
        "Backups use the suffix .backup-YYYYMMDDTHHMMSSZ beside the original file. "
        "The installer also enables UFW, Nginx, and the socketio-server service.",
    )

    add_heading(doc, 1, "5. Alert contract")
    add_table(
        doc,
        ["Direction", "Event", "Payload"],
        [
            ["Edge -> relay", "inspection_status", "JSON alert object"],
            ["Edge -> relay", "alert:publish", "Alternative JSON alert object"],
            ["Relay -> browser", "alert_received", "{ received_at, alert } envelope"],
            ["Relay -> new client", "server:ready", "Socket ID, message, timestamp"],
        ],
        [1900, 2200, 5260],
    )
    add_text(doc, "All clients authenticate during the Socket.IO handshake with auth.token.")
    add_code(
        doc,
        '''{
  "received_at": "2026-07-29T08:13:25.120Z",
  "alert": {
    "event": "night_person_present",
    "site": "Jutemill",
    "camera": "Gate Camera",
    "status": "alert",
    "confidence": 0.93,
    "details": "Person detected at the gate"
  }
}''',
    )

    add_heading(doc, 1, "6. Configure the VisualAI Edge publisher")
    add_text(doc, "Run the project from VisualAI-Edge / Projects / Jutemill / main.py. Keep the token outside source control.")
    add_code(
        doc,
        """SOCKET_IO_SERVER_URL=https://socket.sukalyanai.com
SOCKET_AUTH_TOKEN=<same-token-generated-on-contabo>""",
    )
    add_bullet(doc, "The edge publisher sends inspection_status and waits for the relay acknowledgement.")
    add_bullet(doc, "The current transport test can emit an alert every two seconds for commissioning.")
    add_bullet(doc, "MQTT may run in parallel for testing, but the website table uses Socket.IO.")
    add_bullet(doc, "Start locally: cd Projects/Jutemill, activate .venv, then run python main.py.")
    add_callout(
        doc,
        "EXPECTED LOG",
        "[ALERT SENDER] Connected to remote server followed by [ALERT SENDER] JSON alert sent.",
        fill="EBF7F0",
        accent=GREEN,
    )

    add_heading(doc, 1, "7. Configure the Hostinger frontend")
    add_text(doc, "The factory automation page connects to the public relay URL and listens for alert_received.")
    add_code(
        doc,
        """VITE_SOCKET_IO_URL=https://socket.sukalyanai.com
VITE_SOCKET_IO_TOKEN=<same-token-generated-on-contabo>""",
    )
    add_bullet(doc, "Store the token as a protected GitHub Actions secret; do not commit it to the repository.")
    add_bullet(doc, "Pass the secret to the frontend build environment during CI/CD.")
    add_bullet(doc, "After deployment, hard-refresh https://www.sukalyanai.com/products/factory-automation.")
    add_bullet(doc, "The status should read Alerts: Connected; new alerts appear at the top of the vertically scrolling table.")
    add_callout(
        doc,
        "BROWSER SECURITY NOTE",
        "A token compiled into frontend JavaScript can be viewed by visitors. This shared token is acceptable only "
        "for controlled testing. Production should use short-lived viewer tokens or server-side authorization.",
        fill="FDECEC",
        accent=RED,
    )

    add_heading(doc, 1, "8. Acceptance test")
    add_step(doc, "Check the public relay.", "Open the health endpoint and confirm status is ok.")
    add_code(doc, "curl https://socket.sukalyanai.com/health")
    add_step(doc, "Check the VPS service.", "Confirm it is active and inspect live connection logs.")
    add_code(
        doc,
        """sudo systemctl status socketio-server --no-pager
sudo journalctl -u socketio-server -f""",
    )
    add_step(doc, "Open the website.", "Confirm Alerts: Connected before starting the edge process.")
    add_step(doc, "Start Jutemill main.py.", "Observe JSON alert sent at the configured interval.")
    add_step(doc, "Verify delivery.", "The received count increments and the newest row remains first.")
    add_table(
        doc,
        ["Check", "Pass condition"],
        [
            ["Health endpoint", 'HTTP 200 and "service":"socketio-alert-relay"'],
            ["VPS journal", "Publisher and browser client connection entries appear"],
            ["Edge terminal", "Connected to remote server; JSON alert sent"],
            ["Website badge", "Alerts: Connected"],
            ["Website table", "Count increments; newest alert is at the top"],
        ],
        [2800, 6560],
    )

    add_heading(doc, 1, "9. Routine operations")
    add_table(
        doc,
        ["Task", "Command"],
        [
            ["Service status", "sudo systemctl status socketio-server --no-pager"],
            ["Live logs", "sudo journalctl -u socketio-server -f"],
            ["Restart relay", "sudo systemctl restart socketio-server"],
            ["Validate Nginx", "sudo nginx -t"],
            ["Reload Nginx", "sudo systemctl reload nginx"],
            ["Certificate test", "sudo certbot renew --dry-run"],
            ["Show config keys only", "sudo awk -F= '/^[A-Z_]+=/{print $1}' /etc/socketio-server.env"],
        ],
        [2600, 6760],
    )
    add_text(doc, "Never paste /etc/socketio-server.env into tickets or chat because it contains the authentication token.")

    add_heading(doc, 2, "Token rotation")
    add_step(doc, "Rotate on Contabo.", "Rerun the installer with --rotate-token and save the newly printed value.")
    add_step(doc, "Update the edge secret.", "Replace SOCKET_AUTH_TOKEN and restart main.py.")
    add_step(doc, "Update the CI/CD secret.", "Rebuild and redeploy the Hostinger frontend.")
    add_step(doc, "Confirm both clients reconnect.", "Review the website badge and the VPS journal.")
    add_callout(
        doc,
        "EXPECTED INTERRUPTION",
        "Rotation immediately disconnects clients using the old token. Update the edge and frontend together during a maintenance window.",
        fill="FFF6E5",
        accent=AMBER,
    )

    add_heading(doc, 1, "10. Troubleshooting")
    add_table(
        doc,
        ["Symptom", "Likely cause", "Action"],
        [
            ["Website says Disconnected", "DNS, TLS, origin, or token mismatch", "Check /health, browser console, ALLOWED_ORIGINS, and the CI/CD token."],
            ["Connected but no alert rows", "Publisher event or relay version mismatch", "Tail the VPS journal; confirm edge emits inspection_status and browser listens for alert_received."],
            ["Service restart loop", "Missing dependency or environment token", "Run journalctl -u socketio-server -n 100; verify node_modules and SOCKET_AUTH_TOKEN."],
            ["502 Bad Gateway", "Relay is stopped or wrong local port", "Check service status and ensure Nginx proxies to 127.0.0.1:3000."],
            ["Unauthorized connection", "Different tokens on client and VPS", "Update secrets from the same generated token; restart/redeploy clients."],
            ["CORS error", "Website origin absent", "Set both https://www.sukalyanai.com and https://sukalyanai.com, then rerun installer."],
            ["Certificate failure", "DNS not pointing to VPS or ports blocked", "Correct the A record, allow ports 80/443, then rerun without --skip-certbot."],
        ],
        [2350, 2850, 4160],
    )

    add_heading(doc, 1, "11. Recovery and archive controls")
    add_bullet(doc, "Installer backups are timestamped next to each replaced configuration file.")
    add_bullet(doc, "To restore, stop the service, copy the chosen backup over the active file, restore ownership, and restart.")
    add_bullet(doc, "Keep the installer and this runbook in source control; keep tokens only in the VPS environment and secret stores.")
    add_bullet(doc, "Record each production token rotation date, operator, and deployment commit in the change log.")
    add_bullet(doc, "Do not expose TCP 3000 publicly; Nginx on 443 is the only public application entry point.")

    add_heading(doc, 2, "Change log")
    add_table(
        doc,
        ["Version", "Date", "Change", "Owner"],
        [["1.0", "29 Jul 2026", "Initial archived runbook and standalone installer", "SukalyanAI"]],
        [1200, 1700, 4660, 1800],
    )

    core = doc.core_properties
    core.title = "SukalyanAI Socket.IO Alert Delivery System - Deployment Runbook"
    core.subject = "Contabo installation, VisualAI Edge publishing, and Hostinger alert display"
    core.author = "SukalyanAI"
    core.keywords = "Socket.IO, Contabo, VisualAI Edge, Hostinger, alert relay"
    core.comments = "Generated as an operational archive; secrets intentionally omitted."
    core.created = datetime(2026, 7, 29, 0, 0, 0)
    core.modified = datetime(2026, 7, 29, 0, 0, 0)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
